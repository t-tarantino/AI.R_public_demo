"""FASE 1 — Estrazione a revisione dal 'REGISTRO PALLONCINI.xlsx'.

NON scrive nulla nel database. Legge un foglio-stagione (es. '2024-2025'),
interpreta i blocchi settimanali e produce, in data/import_review/:

  * ricavi_<stagione>.csv  — una riga per voce RICAVI (candidato order_item)
  * costi_<stagione>.csv   — una riga per voce COSTI  (candidato expense_item)
  * report_<stagione>.md   — mappature proposte, voci da rivedere e
                             cross-check dei totali settimana per settimana.

Il foglio e' organizzato a blocchi verticali "SETTIMANA g/m". Ogni blocco ha
colonne parallele: COSTI (C=nome, D=codice fornitore, E=euro) e
RICAVI (F=nome, G=codice composizione, H=euro).

    python -m scripts.extract_registro "2024-2025"
"""
from __future__ import annotations

import csv
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT.parent / "REGISTRO PALLONCINI.xlsx"
OUT_DIR = ROOT / "data" / "import_review"

# Colonne (1-based) dentro un blocco settimanale.
COL_COST_NAME, COL_COST_CODE, COL_COST_EUR = 3, 4, 5      # C D E
COL_REV_NAME, COL_REV_CODE, COL_REV_EUR = 6, 7, 8         # F G H

SKIP_NAMES = {"totale", "percentuale", "nome", "codice"}
SKIP_CODES = {"", "tot", "%", "codice"}

# --- Mappa RICAVI: codice legacy del foglio -> codice composizione master ----
# 'conf' = mappatura sicura ; 'rivedere' = ambigua, da confermare col cliente.
REV_CODE_MAP = {
    "1":   ("002", "conf"),      # Arco Maxi
    "10":  ("003", "conf"),      # Arco Maxi Organico
    "11":  ("001", "conf"),      # Mezzo Arco Organico
    "2":   ("004", "conf"),      # Arco Quadrato
    "22":  ("007", "conf"),      # Mezzo Backdrop
    "36":  ("05M", "conf"),      # Colonne (piene) 1,5m -> semplici 150
    "37":  ("06M", "conf"),      # Colonne organiche 1,5m -> 150
    "42":  ("013", "conf"),      # Centrotavola
    "72":  ("011", "conf"),      # Nuvolette a soffitto (a pezzo)
    "74":  ("015", "conf"),      # Palloncini a elio (a pezzo)
    "P01": ("P01", "conf"),
    "P02": ("P02", "conf"),
    "P03": ("P03", "conf"),
    "P04": ("P04", "conf"),
    # --- da rivedere ---
    "35":  ("05S", "rivedere"),  # "Colonne semplici/base": taglia incerta (100?)
    "55":  ("012", "rivedere"),  # "Nuvola/Nuvoletta": 012 con numero? o 011?
    "99":  ("009", "rivedere"),  # "Backdrop small" ma anche "cerchio piccolo"
    "ELI": ("015", "rivedere"),  # "cuori/palloncini a elio": 015? o 016 grappolo?
    "STR": ("",    "rivedere"),  # "Struttura cerchio": forse NON e' una vendita
}

# --- Mappa COSTI: codice legacy -> (party esistente | DA CREARE) --------------
COST_SUPPLIER_MAP = {
    "AMZ": ("Amazon", "conf"),
    "FMG": ("Festival Magia Giocoleria", "conf"),
    "COP": ("COP (DA CREARE — coperture/numeri)", "rivedere"),
    "ELI": ("ELI (DA CREARE — fornitore elio)", "rivedere"),
    "MRK": ("MRK (DA CREARE — marketing/stampa)", "rivedere"),
    "STR": ("STR (DA CREARE — strutture/ferramenta/TEMU)", "rivedere"),
}

# Categoria spesa proposta per codice fornitore (vocabolario controllato).
COST_CATEGORY_HINT = {
    "AMZ": "materials", "FMG": "materials", "COP": "materials",
    "ELI": "materials", "MRK": "marketing", "STR": "equipment",
}

QTY_RE = re.compile(r"^\s*(\d+)\s*[xX]\s*")          # "10x nuvolette"
QTY_RE2 = re.compile(r"^\s*(\d+)\s+[A-Za-z]")        # "18 elio"


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip()


def parse_qty(name: str) -> int:
    for rx in (QTY_RE, QTY_RE2):
        m = rx.match(name)
        if m:
            return int(m.group(1))
    return 1


def week_date(label: str, season: str) -> str:
    """'SETTIMANA 7/9' + stagione '2024-2025' -> '2024-09-07' (ISO)."""
    m = re.search(r"(\d{1,2})\s*/\s*(\d{1,2})", label)
    if not m:
        return ""
    day, month = int(m.group(1)), int(m.group(2))
    y0, y1 = (int(p) for p in season.split("-"))
    year = y0 if month >= 9 else y1
    return f"{year:04d}-{month:02d}-{day:02d}"


def is_data_cell(name, code, eur) -> bool:
    if not isinstance(eur, (int, float)):        # scarta header '€' e '#DIV/0!'
        return False
    if code is not None and norm(code).lower() in SKIP_CODES:
        return False
    if name is not None and norm(name).lower() in SKIP_NAMES:
        return False
    return name is not None or code is not None


def find_blocks(ws):
    """Ritorna [(start_row, label)] per ogni 'SETTIMANA ...' (esclude ESEMPIO)."""
    blocks = []
    for r in range(1, ws.max_row + 1):
        b = ws.cell(row=r, column=2).value
        if isinstance(b, str) and b.strip().upper().startswith("SETTIMANA"):
            if "ESEMPIO" in b.upper():
                continue
            blocks.append((r, b.strip()))
    return blocks


def extract(season: str):
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[season]
    blocks = find_blocks(ws)
    bounds = [(blocks[i][0], blocks[i][1],
               blocks[i + 1][0] if i + 1 < len(blocks) else ws.max_row + 1)
              for i in range(len(blocks))]

    ricavi, costi, checks = [], [], []
    for start, label, end in bounds:
        date = week_date(label, season)
        rev_tot = cost_tot = None
        rev_sum = cost_sum = 0.0
        for r in range(start, end):
            # totali dichiarati nel foglio (riga 'TOT')
            if norm(ws.cell(row=r, column=COL_COST_CODE).value or "").lower() == "tot":
                v = ws.cell(row=r, column=COL_COST_EUR).value
                if isinstance(v, (int, float)):
                    cost_tot = v
            if norm(ws.cell(row=r, column=COL_REV_CODE).value or "").lower() == "tot":
                v = ws.cell(row=r, column=COL_REV_EUR).value
                if isinstance(v, (int, float)):
                    rev_tot = v

            # RICAVI
            f = ws.cell(row=r, column=COL_REV_NAME).value
            g = ws.cell(row=r, column=COL_REV_CODE).value
            h = ws.cell(row=r, column=COL_REV_EUR).value
            if is_data_cell(f, g, h):
                code = norm(g)
                comp, conf = REV_CODE_MAP.get(code, ("", "NON MAPPATO"))
                ricavi.append({
                    "settimana": label, "data": date, "riga_excel": r,
                    "nome": norm(f) if f is not None else "", "codice_legacy": code,
                    "euro": h, "qta_stimata": parse_qty(norm(f) if f else ""),
                    "composizione_proposta": comp, "confidenza": conf,
                })
                rev_sum += float(h)

            # COSTI
            c = ws.cell(row=r, column=COL_COST_NAME).value
            d = ws.cell(row=r, column=COL_COST_CODE).value
            e = ws.cell(row=r, column=COL_COST_EUR).value
            if is_data_cell(c, d, e):
                code = norm(d)
                supp, conf = COST_SUPPLIER_MAP.get(code, ("", "NON MAPPATO"))
                costi.append({
                    "settimana": label, "data": date, "riga_excel": r,
                    "descrizione": norm(c) if c is not None else "",
                    "codice_legacy": code, "euro": e,
                    "fornitore_proposto": supp,
                    "categoria_proposta": COST_CATEGORY_HINT.get(code, "other"),
                    "confidenza": conf,
                })
                cost_sum += float(e)

        checks.append({
            "settimana": label, "data": date,
            "ricavi_somma_righe": round(rev_sum, 2), "ricavi_tot_foglio": rev_tot,
            "ricavi_ok": rev_tot is None or abs(rev_sum - rev_tot) < 0.01,
            "costi_somma_righe": round(cost_sum, 2), "costi_tot_foglio": cost_tot,
            "costi_ok": cost_tot is None or abs(cost_sum - cost_tot) < 0.01,
        })
    return ricavi, costi, checks


def write_csv(path: Path, rows: list[dict]):
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def write_report(path: Path, season, ricavi, costi, checks):
    lines = [f"# Report estrazione — stagione {season}", ""]
    lines += [
        f"- Voci RICAVI estratte: **{len(ricavi)}** (tot € "
        f"{round(sum(r['euro'] for r in ricavi), 2)})",
        f"- Voci COSTI estratte:  **{len(costi)}** (tot € "
        f"{round(sum(c['euro'] for c in costi), 2)})",
        f"- Settimane: **{len(checks)}**", "",
        "## Cross-check totali (righe estratte vs 'TOT' del foglio)", "",
        "| Settimana | Ricavi righe | Ricavi TOT | OK | Costi righe | Costi TOT | OK |",
        "|---|--:|--:|:-:|--:|--:|:-:|",
    ]
    for c in checks:
        lines.append(
            f"| {c['settimana']} | {c['ricavi_somma_righe']} | "
            f"{c['ricavi_tot_foglio']} | {'✅' if c['ricavi_ok'] else '❌'} | "
            f"{c['costi_somma_righe']} | {c['costi_tot_foglio']} | "
            f"{'✅' if c['costi_ok'] else '❌'} |")

    bad = [c for c in checks if not c["ricavi_ok"] or not c["costi_ok"]]
    lines += ["", f"Settimane con discrepanza: **{len(bad)}**", ""]

    rev_rev = sorted({(r["codice_legacy"], r["composizione_proposta"], r["confidenza"])
                      for r in ricavi if r["confidenza"] != "conf"})
    lines += ["## RICAVI — codici da confermare / non mappati", "",
              "| Codice legacy | Composizione proposta | Stato |", "|---|---|---|"]
    lines += [f"| {a} | {b or '—'} | {c} |" for a, b, c in rev_rev]

    cost_rev = sorted({(c["codice_legacy"], c["fornitore_proposto"], c["confidenza"])
                       for c in costi if c["confidenza"] != "conf"})
    lines += ["", "## COSTI — fornitori da creare / confermare", "",
              "| Codice legacy | Fornitore proposto | Stato |", "|---|---|---|"]
    lines += [f"| {a} | {b or '—'} | {c} |" for a, b, c in cost_rev]

    path.write_text("\n".join(lines), encoding="utf-8")


def main():
    season = sys.argv[1] if len(sys.argv) > 1 else "2024-2025"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ricavi, costi, checks = extract(season)
    write_csv(OUT_DIR / f"ricavi_{season}.csv", ricavi)
    write_csv(OUT_DIR / f"costi_{season}.csv", costi)
    write_report(OUT_DIR / f"report_{season}.md", season, ricavi, costi, checks)
    print(f"Stagione {season}: {len(ricavi)} ricavi, {len(costi)} costi, "
          f"{len(checks)} settimane.")
    print(f"Output in {OUT_DIR}")


if __name__ == "__main__":
    main()
